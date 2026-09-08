from __future__ import annotations

import json
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.dependencies import context as context_dependency
from app.core.security import sha256_file
from app.main import create_app
from app.services.exports import export_file
from app.workers.reporting import write_report_html


def test_export_preserves_bytes_and_existing_files_in_chinese_directory(tmp_path):
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"original output")
    directory = tmp_path / "中文 导出文件夹"
    directory.mkdir()
    existing = directory / "模型报告.xlsx"
    existing.write_bytes(b"user file")
    result = export_file(source, "模型报告.xlsx", sha256_file(source), str(directory))
    assert Path(result["path"]).parent == directory
    assert result["filename"] == "模型报告 (1).xlsx"
    assert Path(result["path"]).read_bytes() == source.read_bytes()
    assert existing.read_bytes() == b"user file"


def test_concurrent_exports_never_overwrite_each_other(tmp_path):
    source = tmp_path / "source.html"
    source.write_text("<h1>合成报告</h1>", encoding="utf-8")
    directory = tmp_path / "output"
    directory.mkdir()

    def save(_):
        return export_file(source, "report.html", sha256_file(source), str(directory))

    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(save, range(6)))
    assert len({result["path"] for result in results}) == 6
    assert all(sha256_file(Path(result["path"])) == sha256_file(source) for result in results)


def test_corrupt_export_source_creates_no_file(tmp_path):
    source = tmp_path / "source"
    source.write_bytes(b"changed")
    output = tmp_path / "output"
    output.mkdir()
    with pytest.raises(ValueError, match="ARTIFACT_CHECKSUM_MISMATCH"):
        export_file(source, "report.html", "0" * 64, str(output))
    assert list(output.iterdir()) == []


def test_copy_failure_removes_only_its_partial_export(tmp_path, monkeypatch):
    source = tmp_path / "source"
    source.write_bytes(b"source")
    output = tmp_path / "output"
    output.mkdir()
    keep = output / "report.html"
    keep.write_bytes(b"keep")

    def fail(original, stream, **kwargs):
        stream.write(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr("app.services.exports.shutil.copyfileobj", fail)
    with pytest.raises(ValueError, match="EXPORT_WRITE_FAILED"):
        export_file(source, "report.html", sha256_file(source), str(output))
    assert list(output.iterdir()) == [keep]
    assert keep.read_bytes() == b"keep"


def test_export_api_and_independent_excel_html(golden, app_paths, tmp_path):
    app = create_app(app_paths, auto_migrate=False)
    app.dependency_overrides[context_dependency] = lambda: golden["context"]
    output = tmp_path / "指定文件夹"
    output.mkdir()
    with TestClient(app) as client:
        run_id = golden["run"]["id"]
        endpoint = f"/api/v1/reports/{run_id}/export"
        blocked = client.post(
            endpoint,
            json={"format": "excel", "directory": str(output)},
            headers={"origin": "https://example.invalid"},
        )
        assert blocked.status_code == 403
        assert not list(output.iterdir())
        for kind, suffix in [("excel", ".xlsx"), ("html", ".html"), ("model", ".zip")]:
            response = client.post(endpoint, json={"format": kind, "directory": str(output)})
            assert response.status_code == 200, response.text
            result = response.json()
            path = Path(result["path"])
            assert path.parent == output
            assert path.suffix == suffix
            assert sha256_file(path) == result["sha256"]
            assert path.stat().st_size == result["size_bytes"]
            if kind == "excel":
                workbook = load_workbook(path, read_only=True)
                assert len(workbook.sheetnames) >= 5
                workbook.close()
            if kind == "html":
                text = path.read_text(encoding="utf-8")
                assert "风控模型报告" in text
                assert not re.search(r'<(?:script|link)[^>]+(?:src|href)="https?://', text)
        assert (
            client.post(endpoint, json={"format": "excel", "directory": "relative"}).status_code
            == 400
        )
        assert (
            client.post(endpoint, json={"format": "exe", "directory": str(output)}).status_code
            == 422
        )
        assert (
            client.post(
                "/api/v1/reports/missing/export", json={"format": "html", "directory": str(output)}
            ).status_code
            == 404
        )
        csp = client.get("/").headers["content-security-policy"]
        assert "frame-src 'self'" in csp and "script-src 'self'" in csp
        preview = client.get(f"/api/v1/reports/{run_id}/preview")
        assert preview.status_code == 200
        assert "风控模型报告" in preview.text
        assert preview.headers["content-disposition"] == "inline"
        assert preview.headers["x-frame-options"] == "SAMEORIGIN"
        assert "sandbox" in preview.headers["content-security-policy"]
        assert "default-src 'none'" in preview.headers["content-security-policy"]
        assert client.get("/").headers["x-frame-options"] == "DENY"
        assert client.get("/api/v1/reports/missing/preview").status_code == 404


def test_export_folder_picker_cancellation_has_no_side_effect(app_paths, monkeypatch):
    purposes = []

    def cancel(*, purpose):
        purposes.append(purpose)
        return None

    monkeypatch.setattr("app.api.artifacts.pick_workspace_directory", cancel)
    with TestClient(create_app(app_paths, auto_migrate=False)) as client:
        response = client.post("/api/v1/exports/native-picker")
        assert response.json() == {"path": None, "cancelled": True}
    assert purposes == ["export"]


def test_html_embedded_json_is_lossless_and_cannot_close_script(golden, tmp_path):
    ctx, run = golden["context"], golden["run"]
    artifact = ctx.database.list("artifacts", {"run_id": run["id"], "kind": "report_json"})[0]
    report = json.loads(Path(artifact["path"]).read_text(encoding="utf-8"))
    name = '教学 A&B <测试> </script><script>alert("x")</script>'
    report["project"]["name"] = name
    path = write_report_html(report, tmp_path / "standalone.html")
    source = path.read_text(encoding="utf-8")
    embedded = re.search(
        r'<script type="application/json" id="risk-model-report-data">(.*?)</script>', source, re.S
    )
    assert json.loads(embedded.group(1))["project"]["name"] == name
    assert source.count("</script>") == 1
    assert "<script>alert" not in source


def test_missing_project_conversation_returns_not_found(app_paths):
    with TestClient(create_app(app_paths, auto_migrate=False)) as client:
        assert client.get("/api/v1/projects/missing/conversation").status_code == 404
