from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from sklearn.dummy import DummyClassifier

from app.workers.modeling import ModelBundle
from app.workers.io import safe_file_name
from app.workers.reporting import _append_table, _finish_sheet


@pytest.mark.parametrize(
    "target", ["../../escaped", "/absolute/target", "C:\\outside\\target", "风控/标签", "Y" * 350]
)
def test_model_package_stays_inside_run_for_arbitrary_column_names(context, monkeypatch, target):
    project = context.catalog.create_project("文件名边界")
    run = {"id": "run_boundary_001", "project_id": project["id"]}
    frame = pd.DataFrame({"feature": [1, 2, 3, 4], target: [0, 1, 0, 1]})
    estimator = DummyClassifier().fit(frame[["feature"]], frame[target])
    bundle = ModelBundle("test", "dummy", estimator, ["feature"], "uncalibrated", {}, {})
    captured = []

    class Captured(Exception):
        pass

    def capture_package(actual_bundle, contract, destination, dependency_lock):
        captured.append(destination)
        raise Captured()

    monkeypatch.setattr("app.services.artifacts.build_model_package", capture_package)
    with pytest.raises(Captured):
        context.artifacts.write_model_artifacts(
            run, {"id": "target_001", "target_column": target}, bundle, frame
        )
    directory = context.artifacts.run_dir(project["id"], run["id"])
    assert captured[0].resolve().parent == directory.resolve()
    assert len(captured[0].name.encode("utf-8")) <= 240
    assert "\\" not in captured[0].name


def test_report_cells_preserve_formula_shaped_text_without_executing(tmp_path: Path):
    path = tmp_path / "report.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    value = '=HYPERLINK("https://example.invalid", "synthetic")'
    _append_table(sheet, "列名安全测试", ["列名", "值"], [[value, 12], ["=1+1", 0.4]])
    _finish_sheet(sheet)
    workbook.save(path)
    with path.open("rb") as stream:
        loaded = load_workbook(stream, data_only=False)
        assert loaded.active["A3"].value == value
        assert loaded.active["A3"].data_type == "s"
        assert loaded.active["A4"].data_type == "s"
        assert loaded.active["B3"].value == 12
        loaded.close()


@pytest.mark.parametrize("source", ["asset", "dataset"])
def test_changed_local_data_is_rejected_before_reuse(context, tmp_path, source):
    project = context.catalog.create_project("不可变数据边界")
    frame = pd.DataFrame({"x": [1, 2, 3, 4], "Y": [0, 1, 0, 1]})
    if source == "asset":
        staged = tmp_path / "input.csv"
        frame.to_csv(staged, index=False)
        record = context.catalog.register_asset(project["id"], staged, "input.csv")
    else:
        record = context.catalog.create_dataset_version(
            project["id"], frame, "frozen", [], {}, freeze=True
        )
    Path(record["stored_path"]).write_text("x,Y\n999,1\n", encoding="utf-8")
    with pytest.raises(ValueError, match="DATA_FILE_CHECKSUM_MISMATCH"):
        if source == "asset":
            context.catalog.materialize_asset(record["id"])
        else:
            context.catalog.dataset_frame(record["id"])


@pytest.mark.parametrize(
    "name", ["中文" * 120 + ".csv", "x" * 220 + ".xlsx", "CON.csv", "LPT1.xls"]
)
def test_upload_filename_keeps_supported_extension_and_avoids_windows_devices(name):
    result = safe_file_name(name)
    assert Path(result).suffix == Path(name).suffix
    assert len(result.encode("utf-8")) <= 160
    assert result.split(".", 1)[0].upper() not in {"CON", "LPT1"}
